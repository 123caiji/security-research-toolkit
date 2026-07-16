from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document(r'D:\software\wechat\chat\xwechat_files\wxid_jgdd701oagw822_133c\msg\file\2026-07\2.同入党申请人谈话记录（样例）.docx')

print("=== 文档结构分析 ===")
print(f"段落数: {len(doc.paragraphs)}")
print(f"表格数: {len(doc.tables)}")

section = doc.sections[0]
print(f"\n页面设置:")
print(f"  宽度: {section.page_width.cm} cm")
print(f"  高度: {section.page_height.cm} cm")
print(f"  左边界: {section.left_margin.cm} cm")
print(f"  右边界: {section.right_margin.cm} cm")
print(f"  上边界: {section.top_margin.cm} cm")
print(f"  下边界: {section.bottom_margin.cm} cm")

print("\n=== 段落内容和样式 ===")
for i, para in enumerate(doc.paragraphs):
    print(f"\n段落 {i}:")
    print(f"  文本: '{para.text[:100]}...'")
    print(f"  对齐: {para.alignment}")
    print(f"  首行缩进: {para.paragraph_format.first_line_indent}")
    print(f"  行距: {para.paragraph_format.line_spacing}")
    print(f"  段前间距: {para.paragraph_format.space_before}")
    print(f"  段后间距: {para.paragraph_format.space_after}")
    
    for j, run in enumerate(para.runs):
        print(f"  Run {j}:")
        print(f"    文本: '{run.text[:50]}'")
        print(f"    字体: {run.font.name}")
        print(f"    字号: {run.font.size}")
        print(f"    加粗: {run.font.bold}")
        print(f"    颜色: {run.font.color.rgb}")

if doc.tables:
    print("\n=== 表格内容 ===")
    for i, table in enumerate(doc.tables):
        print(f"\n表格 {i}:")
        print(f"  行数: {len(table.rows)}")
        print(f"  列数: {len(table.columns)}")
        for row_idx, row in enumerate(table.rows):
            cells = [cell.text for cell in row.cells]
            print(f"  行 {row_idx}: {cells}")

print("\n=== 样式定义 ===")
for style in doc.styles:
    if style.type in ('PARAGRAPH', 'CHARACTER'):
        print(f"  {style.name} ({style.type})")